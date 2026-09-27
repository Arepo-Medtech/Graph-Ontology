#!/usr/bin/env python3
"""The SNOMED CT and LOINC codes the RCPA binds to the same data element of its structured cancer reports.

The RCPA SPIA anatomical pathology FHIR mapping workbooks (cervical, colorectal, endometrial, ovary / fallopian tube /
primary peritoneal, polypectomy) give each data element of a structured report ("Colorectal.subject",
"Colorectal.macro.polypSummary" ...) a SNOMED CT code and a LOINC code on the same row. The pair says the two codes
stand for the same report element -- not that they are the same concept (an element's SNOMED code is often a finding or
observable, its LOINC code a question or a panel) -- so the graph loads it as a relation (rcpa:same_data_element).

RCPA content is copyright (NEHTA-RCPA Terms of Use): the workbooks are read in place from ~/Documents/ONTOLOGIES/RCPA_<version>
(or RCPA_DIR), and what is derived from them stays in cache/rcpa/ (git-ignored) -- never committed. Needs openpyxl.

    scripts/rcpa_elements.py      # writes cache/rcpa/element_bindings.json
"""
from __future__ import annotations

import glob
import json
import os
import re
import sys
from pathlib import Path

import openpyxl


ONT_ROOT = os.environ.get("ONTOLOGIES", os.path.expanduser("~/Documents/ONTOLOGIES"))
_dirs = sorted(d for d in glob.glob(os.path.join(ONT_ROOT, "RCPA_v*")) if os.path.isdir(d))
SRC = Path(os.environ.get("RCPA_DIR", _dirs[-1] if _dirs else ""))
OUT = Path("cache/rcpa/element_bindings.json")
SCT_RE, LOINC_RE = re.compile(r"\d{6,18}"), re.compile(r"\d{1,7}-\d")


def codes(cell, pattern) -> list[str]:
    """One code, or the alternatives a row offers ('85298-8 or 33725-3'); placeholders (xxxxx-x) give none."""
    parts = [p.strip() for p in re.split(r"\s+or\s+|\s*[,;/]\s*", str(cell or "").strip()) if p.strip()]
    return [p for p in parts if pattern.fullmatch(p)] if parts and all(pattern.fullmatch(p) for p in parts) else []


def main() -> int:
    if not SRC.is_dir():
        print(f"RCPA folder not found: {SRC!s} (set RCPA_DIR)", file=sys.stderr)
        return 1
    rows, per_book, skipped = [], {}, 0
    for f in sorted(SRC.glob("*FHIR mapping*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        sheet = next(w for w in wb.worksheets if w.title.startswith("Data Elements"))
        it = sheet.iter_rows(values_only=True)
        head = [str(c).strip() if c else "" for c in next(it)]
        ei, si, sd, li, ld = (head.index(h) for h in ("Element", "Snomed Code", "Snomed  description", "LOINC Code", "LOINC description"))
        n = 0
        for r in it:
            if not (r[si] and r[li]):
                continue
            sct, lnc = codes(r[si], SCT_RE), codes(r[li], LOINC_RE)
            if not (sct and lnc):
                skipped += 1
                continue
            for s in sct:
                for l_ in lnc:
                    rows.append({"element": str(r[ei]).strip(), "sct": s, "sct_description": r[sd], "loinc": l_, "loinc_description": r[ld],
                                 "alternative": len(sct) * len(lnc) > 1, "workbook": f.name})
                    n += 1
        per_book[f.name] = n
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"_note": "RCPA SPIA anatomical pathology FHIR mappings: SNOMED CT and LOINC codes bound to the same "
                                        "report data element (scripts/rcpa_elements.py). RCPA copyright: derived data, never committed.",
                               "source": str(SRC.name), "per_workbook": per_book, "rows_with_placeholder_codes": skipped, "rows": rows},
                              indent=1) + "\n")
    print(f"{OUT}: {len(rows)} element bindings from {len(per_book)} workbooks ({skipped} rows with placeholder codes left out)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
